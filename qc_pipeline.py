#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
 侧弯支具出厂质检表：自动生成 → 归档 → 打印 一体化程序
================================================================================
流程（对应手动操作步骤）：
  1. 读取《苏州交复工作跟进表》Excel 中"订单跟踪"表的数据
  2. 以空模板《侧弯支具出厂质检报告.docx》为底稿生成质检表
     （自动填写：模型名称、线材批次、打印设备、理论重量、实际重量；
       检测项勾选/质检人/日期因 Excel 中无数据而留空，与示例一致）
  3. 按 "订单编号-客户名.docx" 命名
  4. 归档到 交复侧弯订单管理/<客户名称>/<订单编号>/ 目录（目录不存在会自动创建）
  5. 按指定打印机打印（默认每份 2 张，可改）

依赖：python-docx、openpyxl（pip install python-docx openpyxl）
系统：Windows + 已安装 Microsoft Word（打印功能需要 Word COM）

用法示例：
  # 完整流程：生成 + 归档 + 打印（不传 --names 时会提示手动输入模型名称）
  python qc_pipeline.py --names 聂梓瑜 马睿杰 方馨悦 张梓菡
  python qc_pipeline.py          # 然后按提示输入模型名称，可多行，空行结束

  # 只生成和归档，不打印
  python qc_pipeline.py --names 聂梓瑜 马睿杰 --skip-print

  # 自定义各路径和打印参数
  python qc_pipeline.py --excel "D:/xx/工作跟进表.xlsx" --template "D:/xx/模板.docx" \
      --printer "4F-Generic 26C-7SeriesPCL" --copies 3 --names 张三 李四

  # 查看所有可用打印机
  python qc_pipeline.py --list-printers
================================================================================
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    import docx
except ImportError:
    sys.exit("[错误] 缺少 python-docx，请先执行: pip install python-docx openpyxl")
try:
    import openpyxl
except ImportError:
    sys.exit("[错误] 缺少 openpyxl，请先执行: pip install python-docx openpyxl")

# ---------------------------------------------------------------------------
# 默认路径（按你当前环境配置，可通过命令行参数覆盖）
# ---------------------------------------------------------------------------
DEFAULT_EXCEL = r"D:\坚果云（勿改名）\苏州交复\苏州交复工作跟进表-20260323.xlsx"
DEFAULT_TEMPLATE = r"C:\Users\Administrator\Desktop\质检表\侧弯支具出厂质检报告.docx"
DEFAULT_BASE_DIR = r"D:\坚果云（勿改名）\苏州交复\01 侧弯订单\交复侧弯订单管理"
DEFAULT_PRINTER = "4F-Generic 26C-7SeriesPCL"
DEFAULT_COPIES = 2

# Excel 表头关键字 -> 取数据用的列
HEADER_KEYS = {
    "订单编号": "订单编号",
    "客户名称": "客户名称",
    "数据名称": "数据名称",
    "材料批次号": "材料批次号",
    "机器编号": "机器编号",
    "理论重量": "理论重量",
    "实际重量": "实际重量",
    "订单类型": "订单类型",
}

# 归档根目录下的售后件子文件夹名（订单类型="售后订单"时归档到这里）
AFTER_SALE_DIR = "0售后件"


def find_column(header_cells):
    """根据表头文字（可能含换行/括号说明）返回每列在行中的下标。"""
    col_index = {}
    for idx, h in enumerate(header_cells):
        if h is None:
            continue
        text = str(h).replace("\n", "").replace(" ", "")
        for key in HEADER_KEYS:
            if key in text and key not in col_index:
                col_index[key] = idx
    missing = [k for k in HEADER_KEYS if k not in col_index]
    if missing:
        raise RuntimeError("Excel 表头缺少必要列: " + ", ".join(missing))
    return col_index


def load_orders(excel_path, names, chooser=None):
    """从 Excel '订单跟踪' 表读取指定患者（数据名称）的订单数据。

    chooser: 同名患者有多条订单时的人工选择回调 chooser(name, candidates) -> order。
            为 None 时走命令行交互选择（input）。
    """
    if not os.path.exists(excel_path):
        raise RuntimeError("找不到 Excel 文件: " + excel_path)
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "订单跟踪" not in wb.sheetnames:
        raise RuntimeError("Excel 中没有 '订单跟踪' 工作表，现有: " + str(wb.sheetnames))
    ws = wb["订单跟踪"]

    rows = list(ws.iter_rows())
    if not rows:
        raise RuntimeError("订单跟踪表为空")
    col = find_column([c.value for c in rows[0]])

    want = [str(n).strip() for n in names]
    # 收集每个名字的全部候选订单（同一名字可能出现多行 = 多次订单）
    candidates = {n: [] for n in want}
    for row in rows[1:]:
        name_val = row[col["数据名称"]].value
        if name_val is None:
            continue
        name = str(name_val).strip()
        if name not in candidates:
            continue
        order = {
            "name": name,
            "order_no": str(row[col["订单编号"]].value or "").strip(),
            "customer": str(row[col["客户名称"]].value or "").strip(),
            "batch": str(row[col["材料批次号"]].value or "").strip(),
            "machine": str(row[col["机器编号"]].value or "").strip(),
            "theory": row[col["理论重量"]].value,
            "actual": row[col["实际重量"]].value,
            "is_after_sale": (str(row[col["订单类型"]].value or "").strip()
                              == "售后订单"),
        }
        candidates[name].append(order)

    result = []
    for n in want:
        cands = candidates[n]
        if not cands:
            raise RuntimeError("Excel 中找不到这些数据名称: " + n)
        if len(cands) == 1:
            o = cands[0]
        else:
            print("[提示] 患者 '%s' 在 Excel 中有 %d 条订单，需要人工选择"
                  % (n, len(cands)))
            o = (chooser(n, cands) if chooser else _interactive_choose(n, cands))
        result.append(o)
        print("[读取] %s | 订单 %s | 客户 %s | 批次 %s | 设备 %s | 理论 %s g | 实际 %s g%s"
              % (o["name"], o["order_no"], o["customer"], o["batch"],
                 o["machine"], o["theory"], o["actual"],
                 " | 售后件" if o["is_after_sale"] else ""))
    return result


def _interactive_choose(name, candidates):
    """命令行交互：列出候选订单让用户输入序号选择。"""
    print("患者 '%s' 的订单列表：" % name)
    for i, o in enumerate(candidates, 1):
        print("  %d) 订单 %s | 客户 %s | 批次 %s | 设备 %s | 理论 %s g | 实际 %s g%s"
              % (i, o["order_no"], o["customer"], o["batch"], o["machine"],
                 o["theory"], o["actual"], " | 售后件" if o["is_after_sale"] else ""))
    while True:
        try:
            s = input("请输入要生成质检表的订单序号 (1-%d): " % len(candidates)).strip()
        except EOFError:
            raise RuntimeError("未选择订单（输入结束）")
        if s.isdigit() and 1 <= int(s) <= len(candidates):
            return candidates[int(s) - 1]
        print("输入无效，请输入 1-%d 之间的数字" % len(candidates))


def fill_report(template_path, order):
    """复制模板并填写一份质检表，返回生成的 docx 路径。"""
    if not os.path.exists(template_path):
        raise RuntimeError("找不到模板文件: " + template_path)
    if not order["order_no"]:
        raise RuntimeError("订单 %s 缺少订单编号，无法命名" % order["name"])
    filename = "%s-%s.docx" % (order["order_no"], order["name"])
    # 生成到系统临时目录，归档步骤再移动到目标文件夹
    target = os.path.join(tempfile.gettempdir(), filename)
    shutil.copy2(template_path, target)

    d = docx.Document(target)

    # P1 段：模型名称 / 线材批次 / 打印设备（结构固定：run0、run1=空格、run2）
    p1 = d.paragraphs[1]
    p1.runs[0].text = "模型名称：%s       线材批次：%s" % (order["name"], order["batch"])
    p1.runs[2].text = "打印设备：%s" % order["machine"]

    # 表格第 7 行（下标6）重量检测行：理论重量 / 实际重量
    row6 = d.tables[0].rows[6]
    seen = set()
    for cell in row6.cells:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        txt = cell.text
        if txt.startswith("理论重量：") and order["theory"] is not None:
            cell.paragraphs[0].runs[0].text = "理论重量：%sg" % order["theory"]
        elif txt.startswith("实际重量：") and order["actual"] is not None:
            cell.paragraphs[0].runs[0].text = "实际重量：%sg" % order["actual"]

    d.save(target)
    return target


def archive_report(docx_path, base_dir, order):
    """把质检表移动到 交复侧弯订单管理/<客户名称>/<订单编号>/ 目录（售后订单归档到 0售后件 子目录），返回最终路径。"""
    customer = order["customer"] or "未分类客户"
    if order.get("is_after_sale"):
        folder = os.path.join(base_dir, AFTER_SALE_DIR, customer, order["order_no"])
    else:
        folder = os.path.join(base_dir, customer, order["order_no"])
    os.makedirs(folder, exist_ok=True)
    final = os.path.join(folder, os.path.basename(docx_path))
    if os.path.exists(final):
        os.remove(final)  # 覆盖旧版同名文件（同订单重跑场景）
    shutil.move(docx_path, final)
    return final


def list_printers():
    """列出系统可用打印机（后台静默调用，不弹黑色命令窗口）。"""
    r = subprocess.run(
        ["powershell", "-NoProfile", "-Command",
         "Get-Printer | Select-Object -ExpandProperty Name"],
        capture_output=True,
        creationflags=subprocess.CREATE_NO_WINDOW)
    out = r.stdout.decode("gbk", errors="replace") if r.stdout else ""
    printers = [ln.strip() for ln in out.splitlines() if ln.strip()]
    return printers


def print_reports(paths, printer, copies, auto_yes=False):
    """用 Word COM 逐份打印，每份 copies 张。中文路径走 UTF-8 文件，ps1 保持纯 ASCII。"""
    if not paths:
        print("[提示] 没有需要打印的文件")
        return
    if not auto_yes:
        ans = input("将打印 %d 个文件（各 %d 张）到打印机 '%s'，确认？[y/N] "
                    % (len(paths), copies, printer)).strip().lower()
        if ans != "y":
            print("已取消打印")
            return

    # 写路径清单（UTF-8 无 BOM）：第一行打印机名，其余为文件路径
    fd, pathfile = tempfile.mkstemp(suffix=".txt")
    with os.fdopen(fd, "w", encoding="utf-8") as fp:
        fp.write(printer + "\n")
        for p in paths:
            fp.write(p + "\n")

    ps1 = r'''
$lines = Get-Content -Path $args[0] -Encoding UTF8
$printer = $lines[0].Trim()
$paths = $lines[1..($lines.Length - 1)] | Where-Object { $_.Trim().Length -gt 0 }
foreach ($f in $paths) {
    $ok = $false
    for ($try = 1; $try -le 3 -and -not $ok; $try++) {
        $word = $null
        $doc = $null
        try {
            $word = New-Object -ComObject Word.Application
            $word.Visible = $false
            $word.DisplayAlerts = 0
            $word.ActivePrinter = $printer
            $doc = $word.Documents.Open($f)
            $doc.PrintOut($false, $false, 0, '', '', '', 0, %d, '', 1, $false, $true)
            # wait until the background print job is fully submitted,
            # otherwise quitting right away raises RPC error 0x800706BE
            for ($w = 0; $w -lt 240; $w++) {
                try { $pending = $word.BackgroundPrintingStatus } catch { $pending = 0 }
                if ($pending -le 0) { break }
                Start-Sleep -Milliseconds 500
            }
            $doc.Close(0)
            $doc = $null
            $ok = $true
            Write-Output ("printed x%d: " + $f)
        } catch {
            Write-Output ("try" + $try + " failed: " + $f + " | " + $_.Exception.Message)
            Start-Sleep -Seconds 2
        } finally {
            if ($doc -ne $null) { try { $doc.Close(0) } catch {} }
            if ($word -ne $null) {
                # print is already submitted; Quit failing must NOT retry the file,
                # otherwise the same report gets printed twice
                try { $word.Quit() } catch {
                    Write-Output ("quit warning (ignored): " + $_.Exception.Message)
                }
                try { [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null } catch {}
            }
        }
    }
    if (-not $ok) { Write-Output ("GIVEUP: " + $f) }
}
''' % (copies, copies)

    # ps1 必须纯 ASCII，避免 PowerShell 5.1 按 ANSI 读取导致乱码
    fd2, ps1_path = tempfile.mkstemp(suffix=".ps1")
    with os.fdopen(fd2, "w", encoding="ascii") as fp:
        fp.write(ps1)
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", ps1_path, pathfile],
            capture_output=True,
            creationflags=subprocess.CREATE_NO_WINDOW)
        out = r.stdout.decode("gbk", errors="replace") if r.stdout else ""
        err = r.stderr.decode("gbk", errors="replace") if r.stderr else ""
        for line in out.splitlines():
            if "quit warning (ignored)" in line:
                print("  (Word 清理提示已忽略，打印任务已成功提交)")
            else:
                print(line)
        if r.returncode != 0:
            print(err)
    finally:
        os.remove(pathfile)
        os.remove(ps1_path)


def main():
    ap = argparse.ArgumentParser(
        description="侧弯支具出厂质检表：读Excel → 生成质检表 → 归档 → 打印 一体化工具")
    ap.add_argument("--excel", default=DEFAULT_EXCEL, help="工作跟进表 Excel 路径")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE, help="空质检表模板 docx 路径")
    ap.add_argument("--base-dir", default=DEFAULT_BASE_DIR, help="订单管理根目录（客户文件夹所在）")
    ap.add_argument("--printer", default=DEFAULT_PRINTER, help="打印机名称")
    ap.add_argument("--copies", type=int, default=DEFAULT_COPIES, help="每份打印张数")
    ap.add_argument("--names", nargs="+", default=None,
                    help="模型名称（患者名），可多个，如 --names 聂梓瑜 马睿杰；不传则进入手动输入模式")
    ap.add_argument("--skip-print", action="store_true", help="只生成和归档，不打印")
    ap.add_argument("--yes", action="store_true", help="打印前不询问，直接打印")
    ap.add_argument("--list-printers", action="store_true", help="列出可用打印机并退出")
    args = ap.parse_args()

    if args.list_printers:
        for p in list_printers():
            print(p)
        return

    # 未通过 --names 指定时，交互式输入模型名称
    if not args.names:
        print("未指定 --names，进入手动输入模式。")
        print("请输入模型名称（可多个，用空格 / 逗号 / 顿号分隔；直接回车结束输入）：")
        entered = []
        while True:
            try:
                line = input().strip()
            except EOFError:  # 管道/重定向输入结束时停止
                break
            if not line:
                break
            for part in re.split(r"[,，、\s]+", line):
                if part:
                    entered.append(part)
        if not entered:
            sys.exit("[错误] 未输入任何模型名称")
        args.names = entered

    try:
        orders = load_orders(args.excel, args.names)
    except RuntimeError as e:
        sys.exit("[错误] " + str(e))

    print("\n[1/4] 生成质检表 ...")
    temp_files = []
    for o in orders:
        t = fill_report(args.template, o)
        temp_files.append((t, o))
        print("  已生成: %s" % os.path.basename(t))

    print("[2/4] 归档到订单文件夹 ...")
    final_paths = []
    for t, o in temp_files:
        final = archive_report(t, args.base_dir, o)
        final_paths.append(final)
        print("  已归档: %s" % final)

    if args.skip_print:
        print("[3/4] 已跳过打印（--skip-print）")
        print("[4/4] 完成")
        return

    print("[3/4] 打印前检查打印机 ...")
    printers = list_printers()
    if args.printer not in printers:
        print("[警告] 打印机 '%s' 不在系统列表，可用打印机:" % args.printer)
        for p in printers:
            print("  - " + p)
        raise SystemExit("请用 --printer 指定正确的打印机名称")

    print("[4/4] 开始打印 ...")
    print_reports(final_paths, args.printer, args.copies, auto_yes=args.yes)
    print("全部完成")


if __name__ == "__main__":
    main()
